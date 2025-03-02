import os
import requests
import streamlit as st

try:
    import openai
except ImportError:
    openai = None

PUBLISHED = True
APP_URL = "https://cps-import-bot.streamlit.app/"

APP_TITLE = "CPS Canvas Import"
APP_INTRO = "This micro-app allows you import content to the Canvas LMS"

SYSTEM_PROMPT = "Convert raw content into properly formatted HTML excluding any DOCTYPE or extraneous header lines. Additionally, replace specific placeholders like '[begin content block]' with corresponding HTML elements."

HTML_ELEMENTS = {
    "[begin content block]": '<div class="dp-content-block dp-padding-direction-tblr dp-margin-direction-tblr">',
    "[end content block]": '</div>',
    "[begin heading]": '<h2 class="dp-heading dp-padding-direction-tblr dp-margin-direction-tblr">',
    "[end heading]": '</h2>',
    "[begin subheading]": '<h3 class="dp-padding-direction-tblr dp-margin-direction-tblr">',
    "[end subheading]": '</h3>',
    "[begin paragraph]": '<p>',
    "[end paragraph]": '</p>',
    "[begin list]": '<ul>',
    "[end list]": '</ul>',
    "[begin list item]": '<li>',
    "[end list item]": '</li>',
    "[begin table]": '<table class="table-bordered default-base-style">',
    "[end table]": '</table>',
    "[begin table row]": '<tr>',
    "[end table row]": '</tr>',
    "[begin table cell]": '<td>',
    "[end table cell]": '</td>',
}


def extract_text_from_uploaded_files(files):
    texts = []
    for file in files:
        ext = file.name.split('.')[-1].lower()
        if ext == 'docx':
            from docx import Document
            doc = Document(file)
            full_text = "\n".join([para.text for para in doc.paragraphs])
            texts.append(full_text)
        elif ext == 'pdf':
            from pypdf import PdfReader
            pdf = PdfReader(file)
            text = "".join([page.extract_text() for page in pdf.pages if page.extract_text()])
            texts.append(text)
        else:
            texts.append(file.read().decode('utf-8'))
    return "\n".join(texts)


def replace_placeholders_with_html(content):
    """Replaces predefined placeholders with corresponding HTML elements."""
    for placeholder, html_code in HTML_ELEMENTS.items():
        content = content.replace(placeholder, html_code)
    return content


def get_ai_generated_html(prompt):
    openai_api_key = st.secrets["OPENAI_API_KEY"]  
    if not openai_api_key:
        st.error("Missing OpenAI API Key. Please add it to your Streamlit secrets.")
        return None

    headers = {
        "Authorization": f"Bearer {openai_api_key}",
        "Content-Type": "application/json"
    }
    
    payload = {
        "model": "gpt-4o",
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.3
    }

    response = requests.post("https://api.openai.com/v1/chat/completions", headers=headers, json=payload)
    
    if response.status_code == 200:
        ai_response = response.json()["choices"][0]["message"]["content"].strip("`")
        return replace_placeholders_with_html(ai_response)
    else:
        st.error(f"OpenAI API Error: {response.status_code} - {response.text}")
        return None


def check_if_module_exists(module_name, canvas_domain, course_id, headers):
    url = f"https://{canvas_domain}/api/v1/courses/{course_id}/modules"
    response = requests.get(url, headers=headers)
    if response.status_code == 200:
        modules = response.json()
        for module in modules:
            if module["name"].lower() == module_name.lower():
                return module["id"]
    return None


def check_if_page_exists(page_title, canvas_domain, course_id, headers):
    url = f"https://{canvas_domain}/api/v1/courses/{course_id}/pages/{page_title.replace(' ', '-').lower()}"
    response = requests.get(url, headers=headers)
    return response.status_code == 200


def create_or_update_page(page_title, page_body, canvas_domain, course_id, headers):
    if check_if_page_exists(page_title, canvas_domain, course_id, headers):
        url = f"https://{canvas_domain}/api/v1/courses/{course_id}/pages/{page_title.replace(' ', '-').lower()}"
        payload = {"wiki_page": {"body": page_body, "published": PUBLISHED}}
        response = requests.put(url, headers=headers, json=payload)
    else:
        url = f"https://{canvas_domain}/api/v1/courses/{course_id}/pages"
        payload = {"wiki_page": {"title": page_title, "body": page_body, "published": PUBLISHED}}
        response = requests.post(url, headers=headers, json=payload)
    return response.json()


def main():
    st.set_page_config(page_title="Construct HTML Generator", layout="centered", initial_sidebar_state="expanded")
    
    st.title(APP_TITLE)
    st.markdown(APP_INTRO)
    
    st.header("Step 1: Provide Canvas Page Details")
    
    module_title = st.text_input("Enter the title of your module:")
    page_title = st.text_input("Enter the title of your page:")
    uploaded_files = st.file_uploader("Choose files", type=['docx', 'pdf'], accept_multiple_files=True)
    
    uploaded_text = extract_text_from_uploaded_files(uploaded_files) if uploaded_files else ""
    
    if uploaded_text:
        st.markdown("**Extracted Content:**")
        st.text_area("Extracted Text", uploaded_text, height=300)
    
    st.header("Step 2: Generate HTML")
    if st.button("Generate HTML"):
        if not module_title or not page_title or not uploaded_text:
            st.error("Please provide all inputs (module title, page title, and upload at least one file).")
        else:
            preprocessed_text = replace_placeholders_with_html(uploaded_text)
            prompt = f"Module: {module_title}\nPage Title: {page_title}\nContent: {preprocessed_text}"
            ai_generated_html = get_ai_generated_html(prompt)

            if ai_generated_html:
                st.markdown("### AI-Generated HTML Output:")
                st.text_area("AI Response:", ai_generated_html, height=300)
                st.session_state.ai_generated_html = ai_generated_html
            else:
                st.error("AI failed to generate HTML content.")

if __name__ == "__main__":
    main()
